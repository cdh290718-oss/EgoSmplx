"""Export the design document using the applicant's PDF reference style."""
import argparse
import base64
from html import escape
import json
import os
from pathlib import Path
import re
import sys
import subprocess
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))
from copyright_layout import base_document, element, export_pdfs, inline as base_inline, page_field, set_fonts
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont

ASSETS = ROOT / 'docs/assets'
WIDTH = 453.62


def inline(paragraph, text, size=10.5, bold=False):
    def break_identifier(match):
        return re.sub(r'([/_])', lambda token: token[0] + '\u200b', match[0])
    text = re.sub(r'[A-Za-z][A-Za-z0-9_./-]{24,}', break_identifier, text)
    base_inline(paragraph, text, size, bold)


class Figure:
    """Write the same simple diagram as a high-resolution PNG and editable SVG."""
    def __init__(self, name, width, height, background='white'):
        self.name, self.width, self.height = name, width, height
        self.image = Image.new('RGB', (width, height), background)
        self.draw = ImageDraw.Draw(self.image)
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
                    f'<rect width="100%" height="100%" fill="{background}"/>']
        font_root = Path(os.environ.get('EGOSMPLX_REFERENCE_FONT_DIR', os.environ.get('WINDIR', 'C:/Windows') + '/Fonts'))
        self.font = font_root / 'simsun.ttc'
        if not self.font.is_file():
            raise FileNotFoundError('Set EGOSMPLX_REFERENCE_FONT_DIR to a directory containing simsun.ttc')

    def text(self, x, y, value, size=54):
        font = ImageFont.truetype(str(self.font), size)
        self.draw.text((x, y), value, font=font, fill='black', anchor='mt')
        self.svg.append(f'<text x="{x}" y="{y}" font-family="SimSun,宋体" font-size="{size}" fill="black" text-anchor="middle" dominant-baseline="text-before-edge">{escape(value)}</text>')

    def rect(self, x, y, width, height, color='#506e96', radius=22, fill='white', line_width=4):
        self.draw.rounded_rectangle((x,y,x+width,y+height), radius=radius, fill=fill, outline=color, width=line_width)
        self.svg.append(f'<rect x="{x}" y="{y}" width="{width}" height="{height}" rx="{radius}" fill="{fill}" stroke="{color}" stroke-width="{line_width}"/>')

    def line(self, x1, y1, x2, y2, color='#506e96', width=4):
        self.draw.line((x1,y1,x2,y2), fill=color, width=width)
        self.svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="{color}" stroke-width="{width}"/>')

    def arrow(self, x1, y1, x2, y2):
        self.line(x1,y1,x2,y2)
        points = [(x2,y2),(x2-15,y2-23),(x2+15,y2-23)] if x1 == x2 else [(x2,y2),(x2-23,y2-15),(x2-23,y2+15)]
        self.draw.polygon(points, fill='#506e96')
        self.svg.append('<polygon points="'+' '.join(f'{x},{y}' for x,y in points)+'" fill="#506e96"/>')

    def picture(self, source, box, x, y, width):
        cropped = source.crop(box)
        height = round(cropped.height * width / cropped.width)
        cropped = cropped.resize((width,height), Image.Resampling.LANCZOS)
        self.image.paste(cropped, (x,y))
        import io
        buffer = io.BytesIO()
        cropped.save(buffer, format='PNG')
        encoded = base64.b64encode(buffer.getvalue()).decode('ascii')
        self.svg.append(f'<image x="{x}" y="{y}" width="{width}" height="{height}" href="data:image/png;base64,{encoded}"/>')

    def save(self):
        self.image.save(ASSETS/(self.name+'.png'), dpi=(300,300))
        self.svg.append('</svg>')
        (ASSETS/(self.name+'.svg')).write_text('\n'.join(self.svg)+'\n',encoding='utf-8')


def figures():
    with Image.open(ASSETS/'current_controller_stages.jpg') as stages:
        tile = stages.width // 4
        graph = Figure('reference_architecture',2700,1520,'#fff5e9')
        graph.text(1350,65,'总体架构',96)
        names = ['1. 人体初始预测','2. 身体姿态拟合','3. MANO 辅助身体','4. 腕部局部融合']
        notes = [('EgoSMPLX 初始参数','保留原始朝向和平移','人体参数与网格'),
                 ('Sapiens2 身体关节点','刚体及身体关节校正','两阶段身体结果'),
                 ('手腕、掌根与腕环','可见前臂轮廓约束','增强后的身体参数'),
                 ('鱼眼抬升与腕部连接','局部谐波位移','最终网格与验证')]
        colors = ['#bd8962','#78885e','#6b8ba7','#aa963c']
        for k in range(4):
            x = 50+k*670
            graph.rect(x,240,590,1130,colors[k],60,'#fffaf2',3)
            graph.text(x+295,310,names[k],52)
            graph.picture(stages,(k*tile,34,(k+1)*tile,stages.height),x+18,470,554)
            for j,note in enumerate(notes[k]):
                graph.text(x+295,850+j*130,note,48)
            if k<3:
                graph.arrow(x+603,745,x+655,745)
        graph.save()
        for name,indices,labels in [
            ('reference_body_comparison',(0,1),('（a）原始人体预测','（b）两阶段身体拟合')),
            ('reference_fusion_comparison',(2,3),('（a）MANO 辅助身体','（b）最终融合'))]:
            panel = Image.new('RGB',(1400,475),'white')
            draw = ImageDraw.Draw(panel)
            font = ImageFont.truetype(str(graph.font),34)
            for k,index in enumerate(indices):
                pic=stages.crop((index*tile,34,(index+1)*tile,stages.height))
                pic=pic.resize((680,346),Image.Resampling.LANCZOS)
                panel.paste(pic,(k*710,0))
                draw.text((k*710+340,376),labels[k],font=font,fill='black',anchor='mt')
            panel=panel.crop((0,0,1390,433))
            panel.save(ASSETS/(name+'.png'),dpi=(220,220))
    flow=Figure('reference_flow',2700,1960)
    flow.text(1350,30,'软件流程',92)
    steps = [('1. 图像清单与相机标定','读取每帧图像及对应标定'),
             ('2. 运行环境与输入检查','核对路径、尺寸、摘要及模型'),
             ('3. 自动模型预测或缓存读取','人体初值、身体点、分割与手部'),
             ('4. 刚体及身体关节拟合','先调整整体，再拟合身体姿态'),
             ('5. 手部匹配与腕点轮廓校正','冻结左右匹配及可见轮廓'),
             ('6. MANO 辅助身体优化','约束手腕、掌根及腕环'),
             ('7. 鱼眼抬升与原接缝融合','受限腕环候选和局部谐波位移'),
             ('8. 独立重建与结果检查','保存报告及重载、局部穿插结果')]
    for k,(name,note) in enumerate(steps):
        y=170+k*223
        flow.rect(460,y,1500,172,radius=18,line_width=4)
        flow.text(1210,y+22,name,59)
        flow.text(1210,y+98,note,46)
        if k<7:
            flow.arrow(1210,y+178,1210,y+218)
    for start,end,name in [(0,1,'输入检查'),(2,2,'自动观测'),(3,5,'身体拟合'),(6,6,'局部融合'),(7,7,'结果验证')]:
        y1=170+start*223; y2=170+end*223+172
        flow.line(2050,y1,2090,y1,color='#859ec0',width=3)
        flow.line(2090,y1,2090,y2,color='#859ec0',width=3)
        flow.line(2050,y2,2090,y2,color='#859ec0',width=3)
        flow.text(2370,(y1+y2)//2-28,name,55)
    flow.save()
    guide=Figure('reference_guided',2700,950)
    guide.text(1350,30,'MANO 辅助身体优化',82)
    boxes=[(40,'自动观测',['Sapiens2 身体点','WiLoR 手腕与掌根','MANO 的 16 点腕环','冻结的可见前臂轮廓']),
           (960,'参数更新',['19 个身体关节','整体朝向和平移','固定体型与手指','累计角度、腕距和深度']),
           (1880,'保存与复核',['最多 1400 步','可行候选与早停','CPU 重建及缩步','身体参数、网格和报告'])]
    for k,(x,name,lines) in enumerate(boxes):
        guide.rect(x,205,780,580,radius=18)
        guide.text(x+390,250,name,68)
        for j,line in enumerate(lines):guide.text(x+390,370+j*90,line,49)
        if k<2:guide.arrow(x+790,505,x+900,505)
    guide.text(1350,850,'没有可行的新参数时，保留增强前身体并记录回退。',49)
    guide.save()
    with Image.open(ASSETS/'current_eight_final.jpg') as gallery:
        half=gallery.height//2
        gallery.crop((0,0,gallery.width,half)).save(ASSETS/'reference_results_normal.png',dpi=(220,220))
        gallery.crop((0,half,gallery.width,gallery.height)).save(ASSETS/'reference_results_difficult.png',dpi=(220,220))


def table(document, rows):
    rows=[row for row in rows if not all(re.fullmatch(r'[:\-\s]+',cell or '-') for cell in row)]
    count=len(rows[0])
    grid=document.add_table(rows=0,cols=count)
    grid.alignment=WD_TABLE_ALIGNMENT.CENTER
    grid.autofit=False
    if count==8: weights=[35]+[65/7]*7
    elif rows[0]==['参数','默认值','用途']:weights=[31,23,46]
    elif rows[0]==['统计项目','增强前基线','当前结果','有效点数']:weights=[43,20,20,17]
    elif rows[0]==['环境','Python','PyTorch','主要用途']:weights=[23,15,24,38]
    elif count==3:weights=[25,31,44]
    else:weights=[100/count]*count
    widths=[WIDTH*w/sum(weights) for w in weights]
    for col,width in zip(grid.columns,widths):col.width=Pt(width)
    borders=element(grid._tbl.tblPr,'tblBorders')
    for name in ['top','bottom']:element(borders,name,val='single',sz=12,color='000000')
    for name in ['left','right','insideH','insideV']:element(borders,name,val='nil')
    margins=element(grid._tbl.tblPr,'tblCellMar')
    for name,value in [('top',0),('bottom',0),('left',70),('right',70)]:element(margins,name,w=value,type='dxa')
    for i,row in enumerate(rows):
        cells=grid.add_row().cells
        properties=grid.rows[-1]._tr.get_or_add_trPr()
        element(properties,'cantSplit')
        if i==0:element(properties,'tblHeader')
        for j,cell in enumerate(cells):
            cell.width=Pt(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph=cell.paragraphs[0]
            fmt=paragraph.paragraph_format
            fmt.first_line_indent=Pt(0);fmt.line_spacing=Pt(20.4)
            fmt.space_before=fmt.space_after=Pt(0)
            fmt.keep_with_next=i < len(rows)-1 if len(rows)<=9 else i==0
            numeric = rows[0][j] in ['默认值','上限','Python','PyTorch','增强前基线','当前结果','有效点数']
            paragraph.alignment=1 if i==0 or count==8 and j>0 or numeric else 0
            inline(paragraph,row[j],9,bold=i==0)
            if i==0:
                cell_borders=element(cell._tc.get_or_add_tcPr(),'tcBorders')
                element(cell_borders,'bottom',val='single',sz=6,color='000000')
    after=document.add_paragraph()
    after.paragraph_format.first_line_indent=Pt(0)
    after.paragraph_format.line_spacing=Pt(6)
    after.paragraph_format.keep_with_next=True


def illustration(document, caption, path):
    with Image.open(path) as image:
        width,height=image.size
    limit=360 if 'results_' not in path.name else 350
    render_width=min(WIDTH,limit*width/height)
    paragraph=document.add_paragraph()
    paragraph.alignment=1
    fmt=paragraph.paragraph_format
    fmt.first_line_indent=Pt(0);fmt.line_spacing=1
    fmt.space_before=Pt(5);fmt.space_after=Pt(0)
    fmt.keep_with_next=True
    paragraph.add_run().add_picture(str(path),width=Pt(render_width))
    label=document.add_paragraph(style='Caption')
    label.alignment=1
    label.paragraph_format.space_before=Pt(4)
    label.paragraph_format.space_after=Pt(5)
    inline(label,caption)


def build(output):
    metadata=json.loads((ROOT/'software.json').read_text(encoding='utf-8'))
    title=metadata['name_zh']+' '+metadata['registration_version']
    document=base_document(title)
    for part in [document.sections[0].header,document.sections[0].first_page_header]:
        for paragraph in part.paragraphs:
            paragraph.paragraph_format.right_indent=Pt(-39.6)
    section=document.sections[0]
    document.styles['Header'].paragraph_format.tab_stops.clear_all()
    header=section.header.paragraphs[0]
    header.clear()
    header.alignment=0
    tabs=header.paragraph_format.tab_stops
    tabs.clear_all()
    tabs.add_tab_stop(Pt(595.32/2-73.7),WD_TAB_ALIGNMENT.CENTER)
    tabs.add_tab_stop(Pt(WIDTH+39.6),WD_TAB_ALIGNMENT.RIGHT)
    set_fonts(header.add_run('\t'+title+'\t第 '),9)
    page_field(header,9,'Times New Roman')
    set_fonts(header.add_run(' 页'),9)
    for level,size in [(1,14),(2,12),(3,12)]:
        style=document.styles[f'Heading {level}']
        fmt=style.paragraph_format
        fmt.space_before=Pt(10 if level==1 else 5)
        fmt.space_after=Pt(0)
        style.font.size=Pt(size)
    for name in ['TOC 1','Caption']:
        fmt=document.styles[name].paragraph_format
        fmt.line_spacing=Pt(15.6 if name=='TOC 1' else 20.4)
        fmt.space_before=fmt.space_after=Pt(0)
    cover=document.add_paragraph(style='Title')
    cover.alignment=1
    fmt=cover.paragraph_format
    fmt.space_before=Pt(120);fmt.space_after=Pt(30);fmt.line_spacing=Pt(35)
    set_fonts(cover.add_run('第一视角下视双目鱼眼相机\n估计身体姿态软件 V1.0'),22,east='黑体',bold=True)
    subtitle=document.add_paragraph()
    subtitle.alignment=1;subtitle.paragraph_format.first_line_indent=Pt(0)
    set_fonts(subtitle.add_run('软件设计说明书'),20,east='黑体',bold=True)
    document.add_page_break()
    toc_title=document.add_paragraph(style='Title')
    toc_title.alignment=1
    toc_title.paragraph_format.space_before=Pt(0)
    toc_title.paragraph_format.space_after=Pt(20)
    set_fonts(toc_title.add_run('目 录'),18,east='黑体',bold=True)
    toc=document.add_paragraph()
    toc.paragraph_format.first_line_indent=Pt(0)
    run=toc.add_run()
    element(run._r,'fldChar',fldCharType='begin')
    element(run._r,'instrText').text=' TOC \\o "1-1" \\h \\z \\u '
    element(run._r,'fldChar',fldCharType='separate')
    element(run._r,'t').text='更新目录'
    element(run._r,'fldChar',fldCharType='end')
    document.add_page_break()
    lines=(ROOT/'docs/DESIGN_zh.md').read_text(encoding='utf-8').splitlines()
    index=0;started=False
    while index<len(lines):
        line=lines[index].strip();index+=1
        if line.startswith('## '):started=True
        if not line or not started:continue
        heading=re.match(r'^(#{2,4})\s+(.+)',line)
        if heading:
            level=len(heading[1])-1
            paragraph=document.add_paragraph(style=f'Heading {level}')
            set_fonts(paragraph.add_run(heading[2]),14 if level==1 else 12,east='黑体',bold=True)
        elif line.startswith('|'):
            rows=[[x.strip() for x in line.strip('|').split('|')]]
            while index<len(lines) and lines[index].strip().startswith('|'):
                rows.append([x.strip() for x in lines[index].strip().strip('|').split('|')]);index+=1
            table(document,rows)
        elif re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',line):
            match=re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',line)
            illustration(document,match[1],ROOT/'docs'/match[2])
        elif line.startswith('```'):
            language=line[3:];block=[]
            while index<len(lines) and not lines[index].strip().startswith('```'):
                block.append(lines[index]);index+=1
            index+=1
            paragraph=document.add_paragraph()
            fmt=paragraph.paragraph_format
            fmt.first_line_indent=Pt(0);fmt.line_spacing=Pt(15 if language!='shell' else 20.4)
            fmt.space_before=Pt(3);fmt.space_after=Pt(5)
            if language in ['json','shell']:
                border=element(paragraph._p.get_or_add_pPr(),'pBdr')
                for side in ['top','left','bottom','right']:element(border,side,val='single',sz=6,space=4,color='000000')
            set_fonts(paragraph.add_run('\n'.join(block)),9.5 if language!='shell' else 10.5,
                      latin='Courier New' if language!='shell' else 'Times New Roman',bold=language=='shell')
        else:
            paragraph=document.add_paragraph()
            next_line=next((value.strip() for value in lines[index:] if value.strip()),'')
            if re.match(r'^表\s*\d+\s',line) and next_line.startswith('|'):
                paragraph.alignment=1
                paragraph.paragraph_format.first_line_indent=Pt(0)
                paragraph.paragraph_format.keep_with_next=True
                paragraph.paragraph_format.space_before=Pt(5)
                paragraph.paragraph_format.space_after=Pt(2)
            else:paragraph.alignment=3
            inline(paragraph,line)
    element(document.settings._element,'updateFields',val='true')
    document.save(output)


def export_design(path):
    if sys.platform != 'win32':
        export_pdfs([path])
        return
    script = r'''param([string]$TaskPath)
$ErrorActionPreference = 'Stop'
$taskWord = $null
$taskDocument = $null
try {
    $taskWord = New-Object -ComObject Word.Application
    $taskWord.Visible = $false
    $taskWord.DisplayAlerts = 0
    $taskWord.AutomationSecurity = 3
    $taskDocument = $taskWord.Documents.Open($TaskPath, $false, $false)
    $taskDocument.Fields.Update() | Out-Null
    foreach ($taskToc in $taskDocument.TablesOfContents) {
        $taskToc.Update()
        $taskToc.Range.ParagraphFormat.LineSpacingRule = 4
        $taskToc.Range.ParagraphFormat.LineSpacing = 15.6
        $taskToc.Range.ParagraphFormat.SpaceBefore = 0
        $taskToc.Range.ParagraphFormat.SpaceAfter = 0
    }
    $taskDocument.Repaginate()
    foreach ($taskToc in $taskDocument.TablesOfContents) { $taskToc.UpdatePageNumbers() }
    $taskDocument.Save()
    $taskDocument.ExportAsFixedFormat([System.IO.Path]::ChangeExtension($TaskPath, '.pdf'), 17)
} finally {
    if ($null -ne $taskDocument) { $taskDocument.Close(0) }
    if ($null -ne $taskWord) { $taskWord.Quit() }
}
'''
    with tempfile.TemporaryDirectory(prefix='egosmplx_design_export_') as temporary:
        script_path = Path(temporary) / 'export.ps1'
        script_path.write_text(script, encoding='utf-8-sig')
        subprocess.run(['powershell.exe', '-NoProfile', '-NonInteractive', '-ExecutionPolicy',
                        'Bypass', '-File', str(script_path), str(path.resolve())], check=True)
    if not path.with_suffix('.pdf').is_file():
        raise RuntimeError('Word did not produce the design PDF')


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--docx-only',action='store_true')
    parser.add_argument('--skip-figures',action='store_true')
    args=parser.parse_args()
    if not args.skip_figures:figures()
    path=ROOT/'copyright/generated/02_软件设计说明书.docx'
    build(path)
    if not args.docx_only:export_design(path)
    print(path)


if __name__=='__main__':main()
