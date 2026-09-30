"""Reference typography, source provenance and DOCX/PDF export for EgoSmplx."""
import argparse
import ast
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from functools import lru_cache

ROOT = Path(__file__).resolve().parents[1]
SOURCE_WIDTH = (11907 - 2 * 1797) / 20
SOURCE_SIZE = 9.5
SOURCE_LINES = 50


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def font_setup(font_path=None):
    """Use reference font metrics without copying proprietary fonts into the repo."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fonts = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts'
    candidates = [font_path] if font_path else [fonts/'simsun.ttc',
        Path('/usr/share/fonts/truetype/arphic-gbsn00lp/gbsn00lp.ttf'),
        Path('/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf')]
    chinese = next((p for p in candidates if p and p.is_file()), None)
    if chinese is None:
        raise FileNotFoundError('Provide --font-path with a Chinese-capable TTF/TTC font')
    pdfmetrics.registerFont(TTFont('EgoChinese', str(chinese), subfontIndex=0))
    latin = fonts/'times.ttf'
    if not latin.is_file():
        latin = Path('/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf')
    if latin.is_file():
        pdfmetrics.registerFont(TTFont('EgoLatin', str(latin)))


@lru_cache(maxsize=4096)
def character_width(char):
    from reportlab.pdfbase import pdfmetrics
    latin = 'EgoLatin' if 'EgoLatin' in pdfmetrics.getRegisteredFontNames() else 'Times-Roman'
    return pdfmetrics.stringWidth(char, 'EgoChinese' if ord(char) >= 0x2E80 else latin,
                                  SOURCE_SIZE)


def wrap_source(text):
    """Wrap the printed copy only, retaining the repository source indentation."""
    parts, line, width = [], '', 0
    for char in text.expandtabs(4):
        char_width = character_width(char)
        if line and width + char_width > SOURCE_WIDTH - 4:
            parts.append(line)
            line, width = '', 0
        line += char
        width += char_width
    return parts + [line]


def source_listing():
    preferred = ['run_example.sh', 'egosmplx_pipeline/__init__.py',
        'egosmplx_pipeline/__main__.py', 'egosmplx_pipeline/configuration.py',
        'egosmplx_pipeline/pipeline.py', 'egosmplx_pipeline/frame.py',
        'egosmplx_pipeline/bootstrap.py', 'egosmplx_pipeline/io.py',
        'egosmplx_pipeline/sampling.py', 'egosmplx_pipeline/egosmplx_predict.py',
        'egosmplx_pipeline/sapiens_pose.py', 'egosmplx_pipeline/sapiens_seg.py',
        'egosmplx_pipeline/wilor_predict.py', 'egosmplx_pipeline/body_common.py',
        'egosmplx_pipeline/body_rigid.py', 'egosmplx_pipeline/body_pose.py',
        'egosmplx_pipeline/camera.py', 'egosmplx_pipeline/projection.py',
        'egosmplx_pipeline/matching.py', 'egosmplx_pipeline/lift.py',
        'egosmplx_pipeline/topology.py', 'egosmplx_pipeline/fusion.py',
        'egosmplx_pipeline/render.py', 'egosmplx_pipeline/report.py']
    discovered = sorted(p.relative_to(ROOT).as_posix() for folder in
        ['egosmplx_pipeline', 'tools', 'tests'] for p in (ROOT/folder).rglob('*.py'))
    order = preferred + [p for p in discovered if p not in preferred]
    inventory, listing = [], []
    for number, name in enumerate(order, 1):
        path = ROOT/name
        text = path.read_text(encoding='utf-8')
        lines = text.splitlines()
        entry = dict(file_id='F%02d' % number, path=name, sha256=digest(path),
            physical_lines=len(lines), nonempty_lines=sum(bool(x.strip()) for x in lines))
        inventory.append(entry)
        description = '命令行运行入口，启动配置检查与完整处理流程。'
        if path.suffix == '.py':
            description = (ast.get_docstring(ast.parse(text)) or '项目程序模块。').splitlines()[0]
        for label, value in [('文件名：', name), ('功能：', description)]:
            for n, part in enumerate(wrap_source(label + value)):
                listing.append(dict(text=part, path=name, source_line=None,
                                    label=label if n == 0 else ''))
        for source_line, line in enumerate(lines, 1):
            if line.strip():
                for continuation, part in enumerate(wrap_source(' ' + line)):
                    listing.append(dict(text=part, path=name, source_line=source_line,
                                        continuation=continuation, label=''))
    return inventory, [listing[i:i+SOURCE_LINES] for i in range(0, len(listing), SOURCE_LINES)]


def element(parent, tag, **attrs):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node


def set_fonts(run, size=None, east='宋体', latin='Times New Roman', bold=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    run.font.name = latin
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii', 'hAnsi', 'cs']:
        fonts.set(qn('w:' + key), latin)
    fonts.set(qn('w:eastAsia'), east)
    for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + key), None)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def page_field(paragraph, size=11, latin='Cambria'):
    run = paragraph.add_run()
    set_fonts(run, size, latin=latin)
    element(run._r, 'fldChar', fldCharType='begin')
    element(run._r, 'instrText').text = ' PAGE \\* MERGEFORMAT '
    element(run._r, 'fldChar', fldCharType='separate')
    element(run._r, 't').text = '1'
    element(run._r, 'fldChar', fldCharType='end')


def base_document(title, source=False):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.enum.style import WD_STYLE_TYPE
    document = Document()
    section = document.sections[0]
    if source:
        section.page_width, section.page_height = Pt(11907/20), Pt(16839/20)
        section.top_margin, section.bottom_margin = Pt(72), Pt(1247/20)
        section.left_margin = section.right_margin = Pt(1797/20)
        section.header_distance, section.footer_distance = Pt(278/20), Pt(0)
        # Only countBy is explicit in the supplied reference document.
        ln = element(section._sectPr, 'lnNumType', countBy=5)
        section._sectPr.remove(ln)
        section._sectPr.insert_element_before(ln, 'w:pgNumType', 'w:cols', 'w:docGrid')
    else:
        section.page_width, section.page_height = Pt(595.32), Pt(841.92)
        section.top_margin, section.bottom_margin = Pt(72), Pt(56.7)
        section.left_margin, section.right_margin = Pt(73.7), Pt(68)
        section.header_distance, section.footer_distance = Pt(28.35), Pt(0)
    grid = section._sectPr.find(qn('w:docGrid'))
    if grid is not None:
        grid.set(qn('w:linePitch'), '360')
    normal = document.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(SOURCE_SIZE if source else 10.5)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    fmt = normal.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(14 if source else 20.4)
    fmt.widow_control = not source
    element(normal._element.get_or_add_pPr(), 'snapToGrid', val=0)
    if source:
        line_style = document.styles.add_style('line number', WD_STYLE_TYPE.CHARACTER, builtin=True)
        line_style._element.attrib.pop(qn('w:customStyle'), None)
        line_style.base_style = document.styles['Default Paragraph Font']
        line_style.font.name = 'Cambria'
        line_style.font.size = Pt(11)
    else:
        fmt.first_line_indent = Pt(21)
        for name, size in [('Title', 22), ('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 11)]:
            style = document.styles[name]
            for border in list(style._element.xpath('./w:pPr/w:pBdr')):
                border.getparent().remove(border)
            style.font.name, style.font.size, style.font.bold = 'Times New Roman', Pt(size), True
            style.font.color.rgb = RGBColor(0, 0, 0)
            style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体')
            fmt = style.paragraph_format
            fmt.first_line_indent = Pt(0)
            fmt.space_before, fmt.space_after = Pt(12 if name == 'Heading 1' else 8), Pt(4)
            fmt.line_spacing, fmt.keep_with_next = Pt(20.4), True
        for name in ['Caption', 'TOC 1', 'TOC 2', 'TOC 3']:
            style = (document.styles[name] if name in document.styles else
                     document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH))
            style.font.name, style.font.size = 'Times New Roman', Pt(10.5)
            style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
            style.paragraph_format.first_line_indent = Pt(0)
            style.paragraph_format.line_spacing = Pt(18)
    header_style = document.styles['Header']
    header_style.font.name = 'Cambria' if source else 'Times New Roman'
    header_style.font.size = Pt(11 if source else 9)
    header_style.paragraph_format.space_before = header_style.paragraph_format.space_after = Pt(0)
    header_style.paragraph_format.line_spacing = 1
    header_style.paragraph_format.first_line_indent = Pt(0)
    header = section.header.paragraphs[0]
    if source:
        # Two empty header paragraphs are part of the supplied reference layout.
        section.header.add_paragraph(style='Header')
        header = section.header.add_paragraph(style='Header')
    header.alignment = 2
    header.paragraph_format.first_line_indent = Pt(0)
    header.paragraph_format.line_spacing = 1
    header.paragraph_format.space_after = Pt(0)
    set_fonts(header.add_run(title + (' ' * 18 if source else ' ' * 8)), 9,
              latin='宋体' if source else 'Times New Roman')
    if not source:
        set_fonts(header.add_run('第 '), 9)
    page_field(header, 11 if source else 9, 'Cambria' if source else 'Times New Roman')
    if source:
        element(element(header._p.get_or_add_pPr(), 'pBdr'), 'bottom',
                val='single', sz=4, space=1, color='auto')
    else:
        set_fonts(header.add_run(' 页'), 9)
        section.different_first_page_header_footer = True
        first = section.first_page_header.paragraphs[0]
        first.alignment = 2
        set_fonts(first.add_run('第 '), 9)
        page_field(first, 9, 'Times New Roman')
        set_fonts(first.add_run(' 页'), 9)
    section.footer.paragraphs[0].clear()
    document.core_properties.title = title + (' 源程序文档' if source else ' 软件设计说明书')
    document.core_properties.author = 'EgoSmplx'
    document.core_properties.subject = '软件著作权登记鉴别材料'
    return document


def source_docx(path, title, pages):
    from docx.shared import Pt
    document = base_document(title, source=True)
    for page_index, page in enumerate(pages):
        for line_index, row in enumerate(page):
            paragraph = document.add_paragraph()
            fmt = paragraph.paragraph_format
            if page_index and line_index == 0:
                fmt.page_break_before = True
            fmt.keep_together = fmt.keep_with_next = False
            fmt.line_spacing = Pt(14)
            if row['label']:
                label = row['label']
                set_fonts(paragraph.add_run(label), SOURCE_SIZE, bold=True)
                set_fonts(paragraph.add_run(row['text'][len(label):]), SOURCE_SIZE)
            else:
                set_fonts(paragraph.add_run(row['text']), SOURCE_SIZE)
    document.save(path)


def inline(paragraph, text, size=10.5, bold=False):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1', text)
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if not token:
            continue
        strong = token.startswith('**') and token.endswith('**')
        code = token.startswith('`') and token.endswith('`')
        value = token[2:-2] if strong else token[1:-1] if code else token
        if code and len(value) > 24:
            value = re.sub(r'([/_])', lambda m: m[0] + '\u200b', value)
        set_fonts(paragraph.add_run(value), size, bold=bold or strong)


def add_table(document, rows):
    from docx.shared import Pt
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    rows = [r for r in rows if not all(re.fullmatch(r'[:\-\s]+', x or '-') for x in r)]
    if not rows:
        return
    count = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=count)
    table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    weights = [max(6, min(34, max(len(r[c]) if c < len(r) else 0 for r in rows)))
               for c in range(count)]
    if rows[0] == ['条件', '处理', '输出']:
        weights = [35, 35, 30]
    widths = [(595.32 - 73.7 - 68) * v / sum(weights) for v in weights]
    for col, width in zip(table.columns, widths):
        col.width = Pt(width)
    borders = element(table._tbl.tblPr, 'tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        element(borders, side, val='single', sz=4, color='D9D9D9')
    margins = element(table._tbl.tblPr, 'tblCellMar')
    for side, value in [('top', 75), ('bottom', 75), ('left', 80), ('right', 80)]:
        element(margins, side, w=value, type='dxa')
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        element(table.rows[-1]._tr.get_or_add_trPr(), 'cantSplit')
        if index == 0:
            element(table.rows[-1]._tr.get_or_add_trPr(), 'tblHeader')
        for col, cell in enumerate(cells):
            cell.width = Pt(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[col] if col < len(row) else ''
            paragraph = cell.paragraphs[0]
            fmt = paragraph.paragraph_format
            fmt.first_line_indent, fmt.line_spacing = Pt(0), Pt(13)
            fmt.space_before = fmt.space_after = Pt(0)
            paragraph.alignment = 1 if len(text) <= 8 else 0
            inline(paragraph, text.replace('<br>', '\n'), size=9, bold=index == 0)
            element(cell._tc.get_or_add_tcPr(), 'shd',
                    fill='E7EDF4' if index == 0 else 'F7F9FB' if index % 2 == 0 else 'FFFFFF')
    after = document.add_paragraph()
    after.paragraph_format.first_line_indent = Pt(0)
    after.paragraph_format.line_spacing = Pt(5)
    after.paragraph_format.space_after = Pt(3)


def image_block(document, caption, target, markdown_path):
    from docx.shared import Pt
    from PIL import Image
    path = (markdown_path.parent/target).resolve()
    if not path.is_file():
        raise FileNotFoundError('Missing document illustration: ' + str(path))
    with Image.open(path) as source:
        width, height = source.size
    render_width = min(595.32 - 73.7 - 68, 465 * width / height)
    paragraph = document.add_paragraph()
    paragraph.alignment = 1
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(6), Pt(4), 1
    fmt.keep_with_next = True
    picture = paragraph.add_run().add_picture(str(path), width=Pt(render_width))
    picture._inline.docPr.set('descr', caption)
    label = document.add_paragraph(style='Caption')
    label.alignment = 1
    label.paragraph_format.keep_together = True
    label.paragraph_format.space_after = Pt(7)
    inline(label, caption, 10.5)


def design_docx(path, title, markdown_path, cover=True):
    from docx.shared import Pt
    document = base_document(title)
    if cover:
        paragraph = document.add_paragraph(style='Title')
        paragraph.alignment = 1
        fmt = paragraph.paragraph_format
        fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(120), Pt(30), Pt(35)
        name = title.replace('第一视角下视双目鱼眼相机', '第一视角下视双目鱼眼相机\n', 1)
        set_fonts(paragraph.add_run(name), 22, east='黑体', bold=True)
        paragraph = document.add_paragraph()
        paragraph.alignment = 1
        paragraph.paragraph_format.first_line_indent = Pt(0)
        set_fonts(paragraph.add_run('软件设计说明书'), 20, east='黑体', bold=True)
        paragraph = document.add_paragraph()
        paragraph.alignment = 1
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(22)
        set_fonts(paragraph.add_run('概要设计与详细设计'), 12)
        document.add_page_break()
        paragraph = document.add_paragraph('目录', style='Title')
        paragraph.alignment = 1
        paragraph.paragraph_format.space_after = Pt(12)
        run = document.add_paragraph().add_run()
        element(run._r, 'fldChar', fldCharType='begin')
        element(run._r, 'instrText').text = ' TOC \\o "1-2" \\h \\z \\u '
        element(run._r, 'fldChar', fldCharType='separate')
        element(run._r, 't').text = '更新目录'
        element(run._r, 'fldChar', fldCharType='end')
        document.add_page_break()
    lines = markdown_path.read_text(encoding='utf-8').splitlines()
    index = 0
    first_main_heading = True
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line or line == '---' or line.startswith('# ') or line == '软件设计说明书':
            continue
        heading = re.match(r'^(#{2,4})\s+(.+)', line)
        if heading:
            level = len(heading[1]) - 1
            paragraph = document.add_paragraph(style='Heading ' + str(level))
            if level == 1:
                paragraph.paragraph_format.page_break_before = not first_main_heading and not heading[2].startswith('第三部分')
                first_main_heading = False
            inline(paragraph, heading[2], size={1:14, 2:12, 3:11}[level], bold=True)
            for run in paragraph.runs:
                set_fonts(run, east='黑体', bold=True)
            continue
        if line.startswith('|'):
            rows = [[v.strip() for v in line.strip('|').split('|')]]
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append([v.strip() for v in lines[index].strip().strip('|').split('|')])
                index += 1
            add_table(document, rows)
            continue
        matched = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', line)
        if matched:
            image_block(document, matched[1], matched[2], markdown_path)
            continue
        if line.startswith('```'):
            language, block = line[3:].strip(), []
            while index < len(lines) and not lines[index].strip().startswith('```'):
                block.append(lines[index])
                index += 1
            index += 1
            if language == 'mermaid':
                continue
            for value in block:
                paragraph = document.add_paragraph()
                fmt = paragraph.paragraph_format
                fmt.first_line_indent, fmt.line_spacing, fmt.space_after = Pt(0), Pt(12), Pt(0)
                set_fonts(paragraph.add_run(value), 8.5, latin='Consolas')
            document.add_paragraph().paragraph_format.line_spacing = Pt(5)
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = 3
        bullet = re.match(r'^[-*]\s+(.+)', line)
        if bullet:
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.left_indent = Pt(12)
            inline(paragraph, '• ' + bullet[1])
        else:
            inline(paragraph, line)
    element(document.settings._element, 'updateFields', val='true')
    document.save(path)


def export_pdfs(paths, backend='auto'):
    if backend == 'auto':
        backend = 'word' if sys.platform == 'win32' else 'libreoffice'
    if backend == 'word':
        if sys.platform != 'win32':
            raise RuntimeError('Word PDF export requires Windows and installed Microsoft Word')
        script = r'''param([string]$InputJson)
$ErrorActionPreference = 'Stop'
$taskWord = $null
try {
    $taskWord = New-Object -ComObject Word.Application
    $taskWord.Visible = $false
    $taskWord.DisplayAlerts = 0
    $taskWord.AutomationSecurity = 3
    foreach ($taskPath in (Get-Content -LiteralPath $InputJson -Raw -Encoding UTF8 | ConvertFrom-Json)) {
        $taskDocument = $null
        try {
            $taskDocument = $taskWord.Documents.Open([string]$taskPath, $false, $false)
            $taskDocument.Repaginate()
            $taskDocument.Fields.Update() | Out-Null
            foreach ($taskToc in $taskDocument.TablesOfContents) { $taskToc.Update() }
            $taskDocument.Repaginate()
            foreach ($taskToc in $taskDocument.TablesOfContents) { $taskToc.UpdatePageNumbers() }
            $taskDocument.Save()
            $taskPdf = [System.IO.Path]::ChangeExtension([string]$taskPath, '.pdf')
            $taskDocument.ExportAsFixedFormat($taskPdf, 17)
            Write-Output ([System.IO.Path]::GetFileName($taskPdf))
        } finally { if ($null -ne $taskDocument) { $taskDocument.Close(0) } }
    }
} finally { if ($null -ne $taskWord) { $taskWord.Quit() } }
'''
        with tempfile.TemporaryDirectory(prefix='egosmplx_doc_export_') as temporary:
            script_path, inputs = Path(temporary)/'export.ps1', Path(temporary)/'inputs.json'
            script_path.write_text(script, encoding='utf-8-sig')
            inputs.write_text(json.dumps([str(p.resolve()) for p in paths], ensure_ascii=False), encoding='utf-8-sig')
            subprocess.run(['powershell.exe', '-NoProfile', '-NonInteractive', '-ExecutionPolicy',
                            'Bypass', '-File', str(script_path), str(inputs)], check=True)
    else:
        executable = shutil.which('soffice') or shutil.which('libreoffice')
        if not executable:
            raise FileNotFoundError('LibreOffice is required for PDF export; --docx-only retains editable source')
        with tempfile.TemporaryDirectory(prefix='egosmplx_lo_') as temporary:
            for path in paths:
                subprocess.run([executable, '-env:UserInstallation=' + Path(temporary).as_uri(),
                    '--headless', '--convert-to', 'pdf', '--outdir', str(path.parent), str(path)], check=True)
    for path in paths:
        if not path.with_suffix('.pdf').is_file():
            raise RuntimeError('PDF export did not produce ' + str(path.with_suffix('.pdf')))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT/'copyright/generated')
    parser.add_argument('--font-path', type=Path)
    parser.add_argument('--only', choices=['requested', 'all'], default='requested')
    parser.add_argument('--pdf-backend', choices=['auto', 'word', 'libreoffice'], default='auto')
    parser.add_argument('--docx-only', action='store_true')
    args = parser.parse_args()
    font_setup(args.font_path)
    args.output.mkdir(parents=True, exist_ok=True)
    metadata = json.loads((ROOT/'software.json').read_text(encoding='utf-8'))
    title = metadata['name_zh'] + ' ' + metadata['registration_version']
    inventory, full_pages = source_listing()
    selected = (list(range(len(full_pages))) if len(full_pages) <= 60 else
                list(range(30)) + list(range(len(full_pages)-30, len(full_pages))))
    outputs = []
    for suffix, pages in [('完整', full_pages), ('交存', [full_pages[i] for i in selected])]:
        path = args.output/('03_源程序文档_' + suffix + '.docx')
        source_docx(path, title, pages)
        outputs.append(path)
    design = args.output/'02_软件设计说明书.docx'
    design_docx(design, title, ROOT/'docs/DESIGN_zh.md')
    outputs.append(design)
    if args.only == 'all':
        for name in ['01_登记信息填表稿', '04_权属核对与提交清单']:
            path = args.output/(name+'.docx')
            design_docx(path, title, ROOT/'copyright'/(name+'.md'), cover=False)
            outputs.append(path)
    if not args.docx_only:
        export_pdfs(outputs, args.pdf_backend)
    page_map = [[dict(printed_line=n+1, **{key: row.get(key) for key in
                 ['path', 'source_line', 'continuation']}) for n, row in enumerate(page)]
                for page in full_pages]
    index = dict(software=title, files=inventory,
        physical_source_lines=sum(x['physical_lines'] for x in inventory),
        nonempty_source_lines=sum(x['nonempty_lines'] for x in inventory),
        printed_source_lines=sum(len(p) for p in full_pages), full_source_pages=len(full_pages),
        deposited_original_pages=[i+1 for i in selected], source_lines_per_page=SOURCE_LINES,
        reference_layout=dict(page_twips=[11907,16839], margins_twips=dict(top=1440,right=1797,
            bottom=1247,left=1797,header=278,footer=0), font_latin='Times New Roman',
            font_east_asia='宋体',font_points=9.5,line_spacing_points=14,line_number_every=5),
        listing_rule='Actual nonempty source plus file/function headings; wrap printed copy only; no filler',
        source_page_map=page_map, ownership_status='Applicant must verify rights and third-party boundaries')
    (args.output/'source_inventory.json').write_text(json.dumps(index,indent=2,ensure_ascii=False)+'\n', encoding='utf-8')
    (ROOT/'docs/SOURCE_INDEX.md').write_text('# 源码索引\n\n' + '\n'.join(
        '- %s：`%s`（%d 行，非空 %d 行）' % (x['file_id'],x['path'],x['physical_lines'],x['nonempty_lines'])
        for x in inventory) + '\n', encoding='utf-8')
    print(json.dumps({k:index[k] for k in ['physical_source_lines','nonempty_source_lines',
        'printed_source_lines','full_source_pages']},ensure_ascii=False))


if __name__ == '__main__':
    main()
