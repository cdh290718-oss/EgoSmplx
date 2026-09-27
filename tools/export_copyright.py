#!/usr/bin/env python3
"""Export editable DOCX and paginated PDF from the actual repository sources.

The full source is always retained. The deposit copy uses the first and last
30 pages when the full listing exceeds 60 pages. No unrelated filler code is
generated. Applicant facts are not inferred from the GitHub account.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re

ROOT = Path(__file__).resolve().parents[1]


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def chunks(items, count):
    return [items[i:i+count] for i in range(0, len(items), count)]


def font_setup(font_path):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = [font_path] if font_path else [
        Path('/usr/share/fonts/truetype/arphic-gbsn00lp/gbsn00lp.ttf'),
        Path('/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf'),
    ]
    path = next((p for p in candidates if p and p.is_file()), None)
    if path is None:
        raise FileNotFoundError('Provide --font-path with a Chinese-capable TTF font')
    pdfmetrics.registerFont(TTFont('EgoChinese', str(path)))
    return path


def wrap(text, size, width):
    """Wrap only the printed representation; source files remain unchanged."""
    from reportlab.pdfbase.pdfmetrics import stringWidth
    lines, line = [], ''
    for char in text.expandtabs(4):
        if line and stringWidth(line + char, 'EgoChinese', size) > width:
            lines.append(line)
            line = ''
        line += char
    if line:
        lines.append(line)
    return lines


def source_listing():
    preferred = ['egosmplx_pipeline/__init__.py', 'egosmplx_pipeline/__main__.py',
                 'egosmplx_pipeline/configuration.py', 'egosmplx_pipeline/pipeline.py',
                 'egosmplx_pipeline/frame.py', 'egosmplx_pipeline/bootstrap.py',
                 'egosmplx_pipeline/io.py', 'egosmplx_pipeline/sampling.py',
                 'egosmplx_pipeline/sapiens_pose.py', 'egosmplx_pipeline/sapiens_seg.py',
                 'egosmplx_pipeline/wilor_predict.py', 'egosmplx_pipeline/body_common.py',
                 'egosmplx_pipeline/body_rigid.py', 'egosmplx_pipeline/body_pose.py',
                 'egosmplx_pipeline/projection.py',
                 'egosmplx_pipeline/matching.py', 'egosmplx_pipeline/lift.py',
                 'egosmplx_pipeline/topology.py', 'egosmplx_pipeline/fusion.py',
                 'egosmplx_pipeline/render.py', 'egosmplx_pipeline/report.py']
    discovered = sorted(str(p.relative_to(ROOT)) for folder in ['egosmplx_pipeline', 'tools', 'tests']
                        for p in (ROOT / folder).rglob('*.py'))
    order = preferred + [name for name in discovered if name not in preferred] + ['run_example.sh']
    inventory, listing = [], []
    for number, name in enumerate(order, 1):
        path = ROOT / name
        lines = path.read_text(encoding='utf-8').splitlines()
        entry = dict(file_id='F%02d' % number, path=name, sha256=digest(path),
                     physical_lines=len(lines), nonempty_lines=sum(bool(x.strip()) for x in lines))
        inventory.append(entry)
        for line_number, text in enumerate(lines, 1):
            if not text.strip():
                continue
            # Prefix allows exact correspondence after a long source line wraps.
            prefix = '%s:%04d ' % (entry['file_id'], line_number)
            parts = wrap(text, 9, 451)
            for continuation, part in enumerate(parts):
                label = prefix if not continuation else prefix.rstrip() + '+ '
                listing.append((label + part, name))
    return inventory, chunks(listing, 50)


def clean_markdown(text):
    text = re.sub(r'!\[[^\]]*\]\([^)]*\)', '', text)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1（\2）', text)
    text = text.replace('**', '').replace('`', '')
    return text


def prose_pages(path):
    rows = []
    for line in clean_markdown(path.read_text(encoding='utf-8')).splitlines():
        if not line.strip():
            continue
        line = re.sub(r'^#+\s*', '', line)
        rows.extend((part, str(path.relative_to(ROOT))) for part in wrap(line, 11, 505))
    return chunks(rows, 40)


def export_docx(path, title, subtitle, pages, source=False, original_numbers=None, figures=()):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin, section.bottom_margin = Pt(50), Pt(42)
    section.left_margin, section.right_margin = Pt(42), Pt(42)
    section.header_distance, section.footer_distance = Pt(18), Pt(18)
    normal = document.styles['Normal']
    normal.font.name = 'SimSun'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
    normal.font.size = Pt(9 if source else 11)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = Pt(14.2 if source else 17.3)
    header = section.header.paragraphs[0]
    header.text = title + '  |  ' + subtitle
    for run in header.runs:
        run.font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页')
    for page_index, page in enumerate(pages):
        if page_index:
            document.add_page_break()
        for line, _ in page:
            paragraph = document.add_paragraph()
            paragraph.add_run(line)
    for caption, image in figures:
        document.add_page_break()
        document.add_paragraph(caption)
        document.add_picture(str(image), width=Inches(7.0))
    document.save(path)


def export_pdf(path, title, subtitle, pages, source=False, original_numbers=None, figures=()):
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    canvas = Canvas(str(path), pagesize=A4)
    canvas.setTitle(title + ' ' + subtitle)
    canvas.setAuthor('EgoSmplx project; applicant facts pending verification')
    width, height = A4
    total = len(pages) + len(figures)
    for index, page in enumerate(pages):
        canvas.setFont('EgoChinese', 8)
        canvas.drawString(42, height-28, title)
        canvas.drawString(42, height-41, subtitle)
        label = '第 %d / %d 页' % (index+1, total)
        if original_numbers is not None:
            label += '；完整源码第 %d 页' % original_numbers[index]
        canvas.drawRightString(width-42, 24, label)
        canvas.setFont('EgoChinese', 9 if source else 11)
        spacing = 14.2 if source else 17.3
        for line_index, (line, _) in enumerate(page):
            canvas.drawString(42, height-64-line_index*spacing, line)
        canvas.showPage()
    for number, (caption, image) in enumerate(figures, len(pages)+1):
        canvas.setFont('EgoChinese', 8)
        canvas.drawString(42, height-28, title)
        canvas.setFont('EgoChinese', 11)
        canvas.drawString(42, height-62, caption)
        reader = ImageReader(str(image))
        iw, ih = reader.getSize()
        scale = min((width-84)/iw, (height-130)/ih)
        canvas.drawImage(reader, 42, height-84-ih*scale, iw*scale, ih*scale)
        canvas.setFont('EgoChinese', 8)
        canvas.drawRightString(width-42, 24, '第 %d / %d 页' % (number, total))
        canvas.showPage()
    canvas.save()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT/'copyright/generated')
    parser.add_argument('--font-path', type=Path)
    args = parser.parse_args()
    font_setup(args.font_path)
    args.output.mkdir(parents=True, exist_ok=True)
    metadata = json.loads((ROOT/'software.json').read_text())
    title = metadata['name_zh'] + ' ' + metadata['registration_version']
    inventory, full_pages = source_listing()
    full_numbers = list(range(1, len(full_pages)+1))
    selected = list(range(len(full_pages))) if len(full_pages) <= 60 else list(range(30)) + list(range(len(full_pages)-30, len(full_pages)))
    deposit = [full_pages[i] for i in selected]
    for suffix, pages, numbers in [('完整', full_pages, full_numbers), ('交存', deposit, [i+1 for i in selected])]:
        stem = args.output / ('03_源程序文档_' + suffix)
        export_pdf(stem.with_suffix('.pdf'), title, '源程序文档（' + suffix + '）', pages, True, numbers)
        export_docx(stem.with_suffix('.docx'), title, '源程序文档（' + suffix + '）', pages, True, numbers)
    figures = [('图1 八帧三阶段输出（左：原始；中：身体校正；右：MANO 融合与局部表面调整）', ROOT/'docs/assets/contact_sheet.jpg')]
    documents = [('01_登记信息填表稿', ROOT/'copyright/01_登记信息填表稿.md', []),
                 ('02_软件设计说明书', ROOT/'docs/DESIGN_zh.md', figures),
                 ('04_权属核对与提交清单', ROOT/'copyright/04_权属核对与提交清单.md', [])]
    for name, path, illustrations in documents:
        pages = prose_pages(path)
        export_pdf(args.output/(name+'.pdf'), title, name[3:], pages, figures=illustrations)
        export_docx(args.output/(name+'.docx'), title, name[3:], pages, figures=illustrations)
    index = dict(software=title, files=inventory,
                 physical_source_lines=sum(x['physical_lines'] for x in inventory),
                 nonempty_source_lines=sum(x['nonempty_lines'] for x in inventory),
                 printed_source_lines=sum(len(p) for p in full_pages), full_source_pages=len(full_pages),
                 deposited_original_pages=[i+1 for i in selected], source_lines_per_page=50,
                 listing_rule='Nonempty actual source lines only; long lines wrap for print with file/line references; no filler',
                 ownership_status='Applicant must verify rights and third-party boundaries')
    (args.output/'source_inventory.json').write_text(json.dumps(index,indent=2,ensure_ascii=False)+'\n')
    (ROOT/'docs/SOURCE_INDEX.md').write_text('# 源码索引\n\n' + '\n'.join(
        '- %s：`%s`（%d 行，非空 %d 行）' % (x['file_id'],x['path'],x['physical_lines'],x['nonempty_lines']) for x in inventory) + '\n')
    print(json.dumps({k:index[k] for k in ['physical_source_lines','nonempty_source_lines','printed_source_lines','full_source_pages']},ensure_ascii=False))


if __name__ == '__main__':
    main()
